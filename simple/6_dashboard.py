# 6_dashboard.py — ШАГ 6. Дашборд: один HTML-файл, открывается двойным щелчком, сервер не нужен.
# Вкладка «Триаж»: 200 алертов выборки 1, решение итоговой ML-модели и объяснение выбранной языковой модели;
#   таблица тандемов: качество объяснений (слепая оценка экспертов из прежней версии), время, цена.
# Вкладка «Сравнение моделей»: пять ML-моделей, ML против LLM в одиночку, признаки, качество по сетям, локальный запуск.
# Оформление — simple/src/dashboard_template.html.
# Результат: simple/results/dashboard.html
# Запуск: python simple/6_dashboard.py
import json, re, sys
from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from src.config import OLD, OUT, PROJECT, utf8
from src.metrics import f1_interval

TEMPLATE = Path(__file__).resolve().parent / "src" / "dashboard_template.html"
LLM = [("q", "Qwen 3.5 4B", "Qwen 3.5 4B"), ("o", "gpt-oss-120b", "gpt-oss-120b"),
       ("g", "GPT-6-sol", "GPT-6-sol"), ("c", "Claude Opus 5.5", "Claude")]   # id, подпись, имя файла тандема
LLM_ALONE_KEY = {"Claude Opus 5.5": "LLM сама решает, атака ли", "GPT-6-sol": "LLM сама решает, атака ли",
                 "gpt-oss-120b": "открытая LLM сама решает", "Llama 3.1 8B": "лучшая локальная LLM",
                 "Qwen 3.5 4B": "локальная LLM"}
SUB = {"XGBoost": "градиентный бустинг, 500 деревьев", "LightGBM": "градиентный бустинг, 500 деревьев",
       "RandomForest": "случайный лес, 200 деревьев", "LogisticRegression": "логистическая регрессия",
       "LinearSVM": "линейный SVM"}


def s(v):
    return "" if pd.isna(v) else str(v).strip()


def levels(p, flag):
    """Приоритет по вероятности ML среди отобранных: верхние 15 % — critical, следующие 25 % — high, остальные — medium."""
    out = np.array(["low"] * len(p), dtype=object)
    idx = np.where(flag)[0]
    if len(idx):
        c, h = np.quantile(p[idx], [0.85, 0.60])
        out[idx] = np.select([p[idx] >= c, p[idx] >= h], ["critical", "high"], "medium")
    return out


def expert_scores():
    """Оценки объяснений (2 эксперта, шкала 0–2): прежняя слепая оценка (48 алертов) + оценка новых объяснений
    простой версии (simple/results/оценка_новых, алерты 49–54). Возвращает {модель: (сумма, число оценок)}."""
    R = PROJECT / "reports"
    N = OUT / "оценка_новых"
    key, acc = {}, {}
    for xl, sheet, col in ((R / "judge_v2" / "Оценка_объяснений_слепая_v2.xlsx", "Оценка", 7),
                           (R / "judge_v3" / "Оценка_объяснений_слепая_v3.xlsx", "Оценка", 7)):
        if xl.exists():
            for r in openpyxl.load_workbook(xl)[sheet].iter_rows(min_row=4, values_only=True):
                if len(r) > col and r[1] in ("A", "B", "C") and isinstance(r[0], (int, float)):
                    key[(int(r[0]), r[1])] = str(r[col]).split(" (")[0]
    xl = N / "Оценка_новых_объяснений_слепая.xlsx"
    if xl.exists():
        for r in openpyxl.load_workbook(xl)["Ключ"].iter_rows(min_row=2, values_only=True):
            key[(int(r[0]), r[1])] = {"Claude Opus 5.5": "Claude"}.get(r[2], r[2])
    files = [R / "judge_v2" / "answers_expert1_part_1.txt", R / "judge_v2" / "answers_expert2_part_1.txt",
             R / "judge_v3" / "answers_expert1_v3.txt", R / "judge_v3" / "answers_expert2_v3.txt",
             N / "answers_expert1.txt", N / "answers_expert2.txt"]
    for f in files:
        if f.exists():
            for n, v, a, b, c in re.findall(r"(\d+)\s+([ABCD])\s+([012])\s+([012])\s+([012])", f.read_text(encoding="utf-8")):
                if (int(n), v) in key:
                    m = key[(int(n), v)]
                    sm, cnt = acc.get(m, (0.0, 0))
                    acc[m] = (sm + (int(a) + int(b) + int(c)) / 3, cnt + 1)
    return acc


def old_probas(fname):
    """Вероятность ML, которая стояла в промпте у объяснений прежней версии (как в 4_tandem: новые файлы важнее)."""
    out = {}
    for stage in ("stage_10_final", "stage_13_ensemble"):
        f = OLD / stage / fname
        if f.exists():
            d = pd.read_csv(f)
            for r in d[d["source"] == "ML+LLM"].itertuples():
                out[int(r.i)] = float(r.ml_proba)
    return out


def load_json(p):
    return json.load(open(p, encoding="utf-8")) if p.exists() else {}


def main():
    utf8()
    res = json.load(open(OUT / "ml_results.json", encoding="utf-8"))
    best, R = res["best"], res["results"]
    summ = json.load(open(OUT / "итог.json", encoding="utf-8"))
    s1 = pd.read_csv(OUT / "sample1.csv")
    T = {cid: pd.read_csv(OUT / "tandem" / f"{fn}.csv") for cid, _, fn in LLM}
    p = T["g"]["ml_proba"].values.astype(float)
    flag = (T["g"]["source"] != "ML").values
    pri = levels(p, flag)
    y1 = s1["label"].ne("benign").values

    from src.llm import LLMS as _L
    OP = {cid: old_probas(_L[fn][2]) for cid, _, fn in LLM}
    alerts = []
    for i, r in s1.iterrows():
        a = {"d": s(r.description), "log": s(r.full_log)[:420], "lab": r.label, "grp": r.group, "sc": r.scenario,
             "p": round(float(p[i]), 3), "pri": pri[i], "x": {}}
        for cid, _, _ in LLM:
            e = T[cid].iloc[i]
            if e.source == "ML+LLM":
                a["x"][cid] = {"pri": s(e.priority), "sig": s(e.attack_signs), "w": s(e.what_happened),
                               "y": s(e.why_suspicious), "r": s(e.recommendation)}
                if s(e.get("explanation_from", "")) == "старая версия" and i in OP[cid]:
                    a["x"][cid]["op"] = round(OP[cid][i], 2)
        alerts.append(a)

    acc = expert_scores()
    oe = load_json(PROJECT / "reports" / "judge_gptoss" / "итоги_оценки.json")
    oq = oe.get("модели", {}).get("gpt-oss-120b (открытая модель)", {}).get("общая")
    oqc = oe.get("модели", {}).get("GPT-6-sol (контроль)", {}).get("общая")
    if oq is not None:                        # gpt-oss: отдельная парная оценка (2 эксперта) + оценка новых объяснений
        sm, cnt = acc.get("gpt-oss-120b", (0.0, 0))
        n_old = 2 * oe.get("n_alerts", 0)
        acc["gpt-oss-120b"] = (sm + oq * n_old, cnt + n_old)
    q = {m: round(sm / cnt, 2) for m, (sm, cnt) in acc.items() if cnt}
    print("качество объяснений:", q, {m: c for m, (_, c) in acc.items()})
    gc_ = load_json(OLD / "stage_13_ensemble" / "gpt_cost.json")
    cm = load_json(OLD / "stage_7" / "claude_polza_measure.json")
    oc = load_json(OLD / "stage_13_ensemble" / "gptoss_cost.json")
    rub = {"q": 0, "g": (gc_["spent_rub_v2"] + gc_["spent_rub_new"]) / (gc_["n_v2"] + gc_["n_new"]) if gc_ else None,
           "o": oc["tandem_rub"] / oc["tandem_n"] if oc else None, "c": cm.get("rub_per_alert")}
    tand = {r["LLM в тандеме"]: r for r in summ["тандем"]}
    n_flag = int(flag.sum())
    llm = []
    for cid, name, fn in LLM:
        t = tand[fn]
        lat = cm.get("lat_s") if cid == "c" else t["медиана секунд LLM на алерт"]
        rr = rub[cid] or 0
        llm.append({"id": cid, "name": name, "tandem": f"{best} → {name}", "where": "на компьютере" if cid == "q" else "облако",
                    "q": q.get({"c": "Claude"}.get(cid, fn)), "qstar": cid == "o",
                    "lat": round(float(lat), 1) if lat not in (None, "не измерялось") else 0,
                    "rub": round(rr, 3 if rr < 0.1 else 2), "tot": round(rr * n_flag, 2),
                    "fp_ok": f'{t["из них LLM: «нет признаков атаки»"]} из {t["ложных тревог среди отобранных"]}'})

    bl = res["baseline"]
    lo, hi = f1_interval(y1, np.ones(len(y1), bool))
    b = R[best]
    ml = []
    for n in sorted(R, key=lambda n: -R[n]["stream"]["AP"]):
        r = R[n]
        ml.append({"name": n + (" (итоговая)" if n == best else ""), "sub": SUB.get(n, ""), "f1": r["sample1"]["F1"],
                   "lo": r["sample1"]["F1_low"], "hi": r["sample1"]["F1_high"], "found": r["sample1"]["found"],
                   "fp": r["sample1"]["fp"], "note": "", "hl": n == best, "f2": r["sample2"]["F1"], "sf1": r["stream"]["F1"],
                   "auc": r["stream"]["ROC_AUC"], "ap": r["stream"]["AP"],
                   "ep": f'{r["episodes"]["found"]} из {r["episodes"]["episodes"]}', "ph": r["episodes"]["flagged_per_hour"]})
    vs = [dict(ml[0], sub="машинное обучение, решает за доли миллисекунды")]
    for r in summ["LLM"]:
        if r["модель"] in LLM_ALONE_KEY:
            vs.append({"name": r["модель"], "sub": LLM_ALONE_KEY[r["модель"]], "f1": r["F1"], "lo": r["F1 от"], "hi": r["F1 до"],
                       "found": r["найдено из 87"], "fp": r["ложных"], "note": "", "hl": False})
    abl = [{"name": r["модель"], "main": "основной" in r["модель"] and "+ название" not in r["модель"] and "скрыты" not in r["модель"],
            "f1": r["выборка 1: F1"], "f2": r["выборка 2: F1"], "sf1": r["поток: F1"], "auc": r["поток: ROC-AUC"]}
           for r in summ["признаки"]]
    dist = b["distinguishable"]
    data = {"best": best,
            "kpi": {"n": len(s1), "flag": n_flag, "found": b["sample1"]["found"], "att": b["sample1"]["attacks"],
                    "fp": b["sample1"]["fp"], "ben": int((~y1).sum()), "f1": b["sample1"]["F1"],
                    "lo": b["sample1"]["F1_low"], "hi": b["sample1"]["F1_high"], "base": bl["sample1"]["F1"],
                    "base2": bl["sample2"]["F1"], "s2": b["sample2"]["F1"], "stream": b["stream"]["F1"],
                    "stream_base": bl["stream"]["F1"], "ep": f'{b["episodes"]["found"]} из {b["episodes"]["episodes"]}',
                    "delay": b["episodes"]["median_delay_s"], "per_hour": b["episodes"]["flagged_per_hour"],
                    "per_hour_all": b["episodes"]["alerts_per_hour"], "dist": round(dist["found"] / dist["attacks"] * 100)},
            "base": {"f1": bl["sample1"]["F1"], "lo": lo, "hi": hi},
            "gptoss_eval": {"n": oe.get("n_alerts", 0), "q": oq or 0, "ctrl": oqc or 0},
            "llm": llm, "alerts": alerts, "ml": ml, "vs": vs, "abl": abl, "sc": summ["по сетям"]}
    html = TEMPLATE.read_text(encoding="utf-8").replace("/*DATA*/null", json.dumps(data, ensure_ascii=False, default=float))
    (OUT / "dashboard.html").write_text(html, encoding="utf-8")
    print("готово:", OUT / "dashboard.html", len(html) // 1024, "КБ")


if __name__ == "__main__":
    main()
