# 49_judge_explanations.py
# Слепая оценка объяснений тандема независимой локальной моделью-судьёй.
# Берёт уже обезличенный лист reports\Оценка_объяснений_слепая.xlsx (варианты A/B/C без названий моделей),
# судья видит те же данные, что и человек: текст алерта + одно объяснение, и ставит 0-2 по трём критериям.
# Соответствие вариант -> модель берётся из листа «Ключ» и раскрывается только при подсчёте.
# Судья по умолчанию — aya-expanse:8b (Cohere): другое семейство, не Qwen / GPT / Claude и не наши triage-модели.
# Запуск: python scripts\49_judge_explanations.py [--model aya-expanse:8b] [--limit N]
# Результат: results/stage_7/judge_<model>_per_item.csv и judge_<model>_summary.csv
import argparse
import json
import re
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.common import ollama_generate
from src.config import RESULTS, utf8_console

XLSX = Path(r"D:\Default Project\reports\Оценка_объяснений_слепая.xlsx")
S7 = RESULTS / "stage_7"

JUDGE_SYSTEM = (
    "Ты — независимый аналитик SOC и эксперт по кибербезопасности. Ты оцениваешь качество объяснений к алертам SIEM. "
    "Оценивай строго и беспристрастно, по тексту алерта. Выводи только один JSON-объект, без пояснений вокруг."
)


def judge_prompt(context, explanation):
    return (
        "Оцени объяснение алерта SIEM по трём критериям, каждое от 0 до 2:\n"
        "- actual: насколько верно описано, что произошло;\n"
        "- why: насколько понятна и обоснована причина подозрительности;\n"
        "- action: насколько рекомендация конкретна и применима.\n"
        "0 — неверно/бесполезно, 1 — частично, 2 — верно и по делу.\n\n"
        f"Данные алерта:\n{context}\n\n"
        f"Объяснение:\n{explanation}\n\n"
        'Ответь строго JSON: {"actual": 0, "why": 0, "action": 0, "comment": "одно предложение"}'
    )


def parse_scores(raw):
    out = {"actual": None, "why": None, "action": None, "comment": ""}
    m = re.search(r"\{.*\}", raw or "", re.S)
    if m:
        try:
            d = json.loads(m.group(0))
        except Exception:
            d = {}
        for k in ("actual", "why", "action"):
            v = d.get(k)
            if v is not None:
                mm = re.search(r"[012]", str(v))
                if mm:
                    out[k] = int(mm.group(0))
        if d.get("comment"):
            out["comment"] = str(d["comment"]).strip()[:300]
    for k in ("actual", "why", "action"):
        if out[k] is None:
            mm = re.search(rf'"{k}"\s*:\s*([012])', raw or "", re.I)
            out[k] = int(mm.group(1)) if mm else None
    return out


def load_blind():
    wb = load_workbook(XLSX, data_only=True)
    ws, kw = wb["Оценка"], wb["Ключ"]
    items, ctx = [], ""
    for r in ws.iter_rows(min_row=4, max_row=ws.max_row):
        a, b, c = r[0].value, r[1].value, r[2].value
        if isinstance(a, str) and a.startswith("Алерт "):
            ctx = a
        elif b in ("A", "B", "C") and isinstance(a, int):
            items.append({"n": a, "variant": b, "context": ctx, "explanation": c})
    keymap = {}
    for row in kw.iter_rows(min_row=2, values_only=True):
        n, i, truth, A, B, C = row
        keymap[n] = {"truth": truth, "A": A, "B": B, "C": C}
    return items, keymap


def main():
    utf8_console()
    ap = argparse.ArgumentParser()
    ap.add_argument("--model", default="aya-expanse:8b")
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    items, keymap = load_blind()
    slug = a.model.replace(":", "_").replace("/", "_")
    out = S7 / f"judge_{slug}_per_item.csv"
    print(f"судья: {a.model} | элементов: {len(items)} (30 алертов x 3 варианта)")

    done = {}
    if out.exists():
        prev = pd.read_csv(out)
        done = {(int(r["n"]), r["variant"]): r.to_dict() for _, r in prev.dropna(subset=["actual"]).iterrows()}
        print(f"продолжение: {len(done)} уже оценено")

    rows = []
    for k, it in enumerate(items):
        n, var = it["n"], it["variant"]
        if (n, var) in done:
            rows.append(done[(n, var)])
            continue
        if a.limit and k >= a.limit:
            continue
        t0 = time.perf_counter()
        try:
            raw = ollama_generate(a.model, judge_prompt(it["context"], it["explanation"]),
                                  system=JUDGE_SYSTEM, num_predict=200, num_ctx=4096)
        except Exception as e:
            raw = ""
            print("  ошибка:", e)
        s = parse_scores(raw)
        dt = time.perf_counter() - t0
        rec = {"n": n, "variant": var, "model": keymap[n][var], "truth": keymap[n]["truth"],
               **s, "latency_s": round(dt, 1), "raw": raw[:400]}
        rows.append(rec)
        print(f"[{len(rows):>3}/{len(items)}] алерт {n} вар.{var} ({keymap[n][var][:22]:<22}) "
              f"-> {s['actual']}/{s['why']}/{s['action']} {dt:.0f}s")
        if len(rows) % 5 == 0:
            pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")

    d = pd.DataFrame(rows)
    d["total"] = d[["actual", "why", "action"]].mean(axis=1)
    d.to_csv(out, index=False, encoding="utf-8-sig")

    summ = (d.groupby("model")[["actual", "why", "action", "total"]]
            .mean().round(2).reset_index().sort_values("total", ascending=False))
    summ["n"] = d.groupby("model").size().values
    sm_out = S7 / f"judge_{slug}_summary.csv"
    summ.to_csv(sm_out, index=False, encoding="utf-8-sig")

    print("\n=== Средние оценки судьи", a.model, "(0-2) ===")
    for _, r in summ.iterrows():
        print(f"  {r['model']:<28} всего {r['total']:.2f} | произошло {r['actual']:.2f} | "
              f"почему {r['why']:.2f} | рекомендация {r['action']:.2f}  (n={int(r['n'])})")
    bad = int(d["actual"].isna().sum())
    print(f"\nне распознано ответов судьи: {bad}; файлы: {out.name}, {sm_out.name}")


if __name__ == "__main__":
    main()
