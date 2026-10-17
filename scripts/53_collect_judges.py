# 53_collect_judges.py
# Собирает оценки: ваши (лист «Оценка» в слепом xlsx) и ИИ-судей (reports/judge/answers_<имя>_part_N.txt),
# считает средние по моделям и согласие. Результат: reports/Оценка_объяснений_итоги.xlsx
import re
from collections import defaultdict
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl import Workbook
ROOT = Path(__file__).resolve().parents[1]
ws = openpyxl.load_workbook(ROOT / "reports" / "Оценка_объяснений_слепая.xlsx")["Оценка"]
model_of, human = {}, {}
for r in ws.iter_rows(min_row=4, values_only=True):
    if r[1] in ("A", "B", "C") and isinstance(r[0], (int, float)):
        key = (int(r[0]), r[1]); model_of[key] = r[7]
        vals = [r[3], r[4], r[5]]
        if all(isinstance(v, (int, float)) for v in vals): human[key] = [int(v) for v in vals]
# Судьи, исключённые из расчёта: почти все оценки «2» (Алиса 91 %, GigaChat 88 %), качество они не различают.
# Файлы ответов сохранены для прозрачности, в средние и согласие они не входят.
EXCLUDE = {"alisa", "gigachat", "judge1"}  # judge1 = прежнее имя файла DeepSeek (дубликат answers_deepseek)
judges = defaultdict(dict)
pat = re.compile(r"(\d+)\D{0,3}([ABCАВС])\D{1,6}([012])\D{1,6}([012])\D{1,6}([012])")
cyr = str.maketrans("АВС", "ABC")
for f in sorted((ROOT / "reports" / "judge").glob("answers_*_part_*.txt")):
    name = f.stem.split("_")[1]
    if name in EXCLUDE:
        continue
    for m in pat.finditer(f.read_text(encoding="utf-8")):  # работает и если ответ в одну строку
        judges[name][(int(m.group(1)), m.group(2).translate(cyr))] = [int(m.group(i)) for i in (3, 4, 5)]
sources = {"Вы": human, **judges}
models = sorted(set(model_of.values()))
wb = Workbook(); s = wb.active; s.title = "Средние по моделям"
s.append(["Оценщик", "Оценено вариантов"] + [f"{m}: общая (0–2)" for m in models])
for nm, d in sources.items():
    row = [nm, len(d)]
    for m in models:
        v = [np.mean(x) for k, x in d.items() if model_of.get(k) == m]
        row.append(round(float(np.mean(v)), 2) if v else "")
    s.append(row)
s.append([])
s.append(["Подкритерии (среднее по всем оценщикам)", "", "Что произошло", "Почему подозрительно", "Рекомендация"])
for m in models:
    cols = []
    for c in range(3):
        v = [x[c] for d in sources.values() for k, x in d.items() if model_of.get(k) == m]
        cols.append(round(float(np.mean(v)), 2) if v else "")
    s.append([m, ""] + cols)
a = wb.create_sheet("Согласие")
a.append(["Пара оценщиков", "Общих вариантов", "Точное совпадение, %", "Корреляция Спирмена (общая оценка)"])
def spearman(x, y):
    from scipy.stats import spearmanr
    r = spearmanr(x, y).correlation
    return "" if r != r else round(float(r), 2)
names = list(sources)
for i in range(len(names)):
    for j in range(i + 1, len(names)):
        d1, d2 = sources[names[i]], sources[names[j]]
        common = sorted(set(d1) & set(d2))
        if len(common) < 5: continue
        ex = np.mean([d1[k] == d2[k] for k in common]) * 100
        a.append([f"{names[i]} и {names[j]}", len(common), round(float(ex)), spearman([np.mean(d1[k]) for k in common], [np.mean(d2[k]) for k in common])])
for sh in wb:
    for col in "ABCDE": sh.column_dimensions[col].width = 34
wb.save(ROOT / "reports" / "Оценка_объяснений_итоги.xlsx")
print({k: len(v) for k, v in sources.items()}, "->", "reports/Оценка_объяснений_итоги.xlsx")
